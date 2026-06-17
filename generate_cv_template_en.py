#!/usr/bin/env python3
"""Generate a BLANK, fillable 3-page CV template (.docx) in ENGLISH,
reproducing the two-column design. All content is placeholder text in
[brackets] that the user replaces with their own information."""

import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image, ImageDraw, ImageFont

OUT = "CV_Template_EN.docx"

# ---- Palette ----------------------------------------------------------------
NAVY = RGBColor(0x14, 0x32, 0x55)
NAVY_HEX = "143255"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BLUE = RGBColor(0x2F, 0x5C, 0x99)
BAR_FILL = "FFFFFF"
BAR_EMPTY = "3C5C86"
HEAD_FONT = "Montserrat"
FONT = "Calibri"

LATO = ("/opt/toolchains/.local/share/mise/installs/ruby/3.4.4/lib/ruby/"
        "3.4.0/rdoc/generator/template/darkfish/fonts/Lato-Regular.ttf")


def make_photo_placeholder(path="template_photo.png"):
    if os.path.exists(path):
        return path
    SIZE = 800
    img = Image.new("RGBA", (SIZE, SIZE), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.ellipse([0, 0, SIZE - 1, SIZE - 1], fill=(255, 255, 255, 255))
    pad = 24
    d.ellipse([pad, pad, SIZE - 1 - pad, SIZE - 1 - pad],
              fill=(220, 228, 240, 255))
    cx = SIZE // 2
    d.ellipse([cx - 120, 230, cx + 120, 470], fill=(143, 162, 189, 255))
    d.pieslice([cx - 210, 520, cx + 210, 940], 180, 360,
               fill=(143, 162, 189, 255))
    try:
        font = ImageFont.truetype(LATO, 90)
    except OSError:
        font = ImageFont.load_default()
    txt = "PHOTO"
    bbox = d.textbbox((0, 0), txt, font=font)
    tw = bbox[2] - bbox[0]
    d.text(((SIZE - tw) / 2 - bbox[0], 690), txt, font=font,
           fill=(255, 255, 255, 235))
    img = img.resize((420, 420), Image.LANCZOS)
    img.save(path)
    return path


# ---- Low-level helpers ------------------------------------------------------
def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("start", start), ("bottom", bottom),
                     ("end", end)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_width(cell, inches):
    cell.width = Inches(inches)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(inches * 1440)))
    tcW.set(qn("w:type"), "dxa")


def add_run(p, text, *, size=11, bold=False, italic=False, color=NAVY,
            font=FONT, caps=False):
    run = p.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    if caps:
        c = OxmlElement("w:caps")
        c.set(qn("w:val"), "true")
        rPr.append(c)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    return run


def space(p, before=0, after=0, line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line


def progress_bar(cell, ratio):
    tbl = cell.add_table(rows=1, cols=2)
    tbl.autofit = False
    total = 1.9
    filled = round(total * ratio, 2)
    empty = round(total - filled, 2)
    left, right = tbl.rows[0].cells
    set_cell_width(left, max(filled, 0.05))
    set_cell_width(right, max(empty, 0.05))
    set_cell_bg(left, BAR_FILL)
    set_cell_bg(right, BAR_EMPTY)
    for c in (left, right):
        set_cell_margins(c, top=10, bottom=10)
        c.paragraphs[0].text = ""
        c.paragraphs[0].add_run(" ").font.size = Pt(3)
    return tbl


def no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        borders.append(el)
    tblPr.append(borders)


# ---- Document ---------------------------------------------------------------
doc = Document()
section = doc.sections[0]
section.page_height = Inches(11.69)
section.page_width = Inches(8.27)
for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(section, attr, Inches(0))

outer = doc.add_table(rows=1, cols=2)
outer.alignment = WD_TABLE_ALIGNMENT.LEFT
outer.allow_autofit = False
sidebar = outer.rows[0].cells[0]
main = outer.rows[0].cells[1]
set_cell_width(sidebar, 2.85)
set_cell_width(main, 5.42)
set_cell_bg(sidebar, NAVY_HEX)
set_cell_bg(main, "FFFFFF")
set_cell_margins(sidebar, top=360, start=360, bottom=360, end=300)
set_cell_margins(main, top=520, start=420, bottom=360, end=420)


# ===== SIDEBAR ===============================================================
def sb_heading(text):
    p = sidebar.add_paragraph()
    space(p, before=16, after=8)
    add_run(p, text, size=15, bold=True, color=WHITE, caps=True, font=HEAD_FONT)


def sb_text(text, *, after=4):
    p = sidebar.add_paragraph()
    space(p, before=0, after=after, line=1.05)
    add_run(p, text, size=11, color=WHITE)
    return p


photo_p = sidebar.paragraphs[0]
photo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
space(photo_p, before=0, after=10)
ph = make_photo_placeholder()
photo_p.add_run().add_picture(ph, width=Inches(1.85))

sb_heading("Contact")
sb_text("Phone: [ Phone number ]", after=6)
sb_text("[ email@address.com ]", after=6)
sb_text("[ Address, City ]", after=6)

sb_heading("Languages")
p = sidebar.add_paragraph(); space(p, before=2, after=1)
add_run(p, "[ Language 1 ]", size=11, color=WHITE)
progress_bar(sidebar, 0.8)
p = sidebar.add_paragraph(); space(p, before=4, after=1)
add_run(p, "[ Language 2 ]", size=11, color=WHITE)
progress_bar(sidebar, 0.6)

sb_heading("Skills")
for _ in range(6):
    sb_text("[ Skill ]")

sb_heading("Interests")
for _ in range(3):
    sb_text("[ Interest ]")


# ===== MAIN COLUMN (PAGE 1) ==================================================
p = main.paragraphs[0]
space(p, before=0, after=0)
add_run(p, "[ FIRST NAME LAST NAME ]", size=28, bold=True, color=NAVY,
        font=HEAD_FONT)
p = main.add_paragraph(); space(p, before=2, after=18)
add_run(p, "[ Job Title ]", size=16, color=LIGHT_BLUE, caps=True,
        font=HEAD_FONT)


def section_title(text):
    p = main.add_paragraph()
    space(p, before=10, after=10)
    add_run(p, text, size=18, bold=True, color=NAVY, caps=True, font=HEAD_FONT)


def entry(title, org, date):
    p = main.add_paragraph()
    space(p, before=8, after=0, line=1.0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(5.0),
                                              WD_TAB_ALIGNMENT.RIGHT)
    add_run(p, title, size=12.5, bold=True, color=NAVY, font=HEAD_FONT)
    add_run(p, "\t" + date, size=10, italic=True, color=LIGHT_BLUE)
    p2 = main.add_paragraph(); space(p2, before=0, after=2, line=1.0)
    add_run(p2, org, size=11.5, italic=True, color=LIGHT_BLUE)


def bullet(text):
    p = main.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.28)
    pf.first_line_indent = Inches(-0.18)
    space(p, before=0, after=2, line=1.0)
    add_run(p, "•  ", size=11, color=LIGHT_BLUE)
    add_run(p, text, size=11, color=LIGHT_BLUE)


section_title("Education")
entry("[ Degree / Program ]", "[ School - City ]", "[ Years ]")
entry("[ Degree / Program ]", "[ School - City ]", "[ Years ]")

section_title("Work Experience")
entry("[ Job Title ]", "[ Company - City ]", "[ Dates ]")
bullet("[ Key responsibility or achievement ]")
bullet("[ Key responsibility or achievement ]")
bullet("[ Key responsibility or achievement ]")
entry("[ Job Title ]", "[ Company - City ]", "[ Dates ]")
bullet("[ Key responsibility or achievement ]")
bullet("[ Key responsibility or achievement ]")
entry("[ Job Title ]", "[ Company - City ]", "[ Dates ]")
bullet("[ Key responsibility or achievement ]")
bullet("[ Key responsibility or achievement ]")

no_borders(outer)
for t in main.tables:
    no_borders(t)
for t in sidebar.tables:
    no_borders(t)


# ===== FULL-WIDTH (PAGES 2 & 3) =============================================
LM = 0.7
RT = 7.5


def banner():
    t = doc.add_table(rows=1, cols=1)
    t.allow_autofit = False
    c = t.rows[0].cells[0]
    set_cell_width(c, 8.27)
    set_cell_bg(c, NAVY_HEX)
    set_cell_margins(c, top=170, start=int(LM * 1440), bottom=170, end=500)
    p = c.paragraphs[0]; space(p, before=0, after=0)
    add_run(p, "[ FIRST NAME LAST NAME ]", size=18, bold=True, color=WHITE,
            font=HEAD_FONT)
    p2 = c.add_paragraph(); space(p2, before=0, after=0)
    add_run(p2, "[ Job Title ]", size=10.5, color=WHITE, caps=True,
            font=HEAD_FONT)
    no_borders(t)
    sp = doc.add_paragraph(); space(sp, before=0, after=6)


def d_title(text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(LM); pf.right_indent = Inches(LM)
    space(p, before=12, after=8)
    add_run(p, text, size=16, bold=True, color=NAVY, caps=True, font=HEAD_FONT)


def d_entry(title, org, date):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(LM); pf.right_indent = Inches(LM)
    pf.tab_stops.add_tab_stop(Inches(RT), WD_TAB_ALIGNMENT.RIGHT)
    space(p, before=8, after=0, line=1.0)
    add_run(p, title, size=12.5, bold=True, color=NAVY, font=HEAD_FONT)
    add_run(p, "\t" + date, size=10, italic=True, color=LIGHT_BLUE)
    p2 = doc.add_paragraph(); p2.paragraph_format.left_indent = Inches(LM)
    space(p2, before=0, after=2, line=1.0)
    add_run(p2, org, size=11.5, italic=True, color=LIGHT_BLUE)


def d_bullet(text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(LM + 0.28); pf.right_indent = Inches(LM)
    pf.first_line_indent = Inches(-0.18)
    space(p, before=0, after=2, line=1.0)
    add_run(p, "•  ", size=11, color=LIGHT_BLUE)
    add_run(p, text, size=11, color=LIGHT_BLUE)


def d_para(text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(LM); pf.right_indent = Inches(LM)
    space(p, before=0, after=6, line=1.15)
    add_run(p, text, size=11, color=NAVY)


def d_label(label, value):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(LM); pf.right_indent = Inches(LM)
    space(p, before=0, after=3, line=1.1)
    add_run(p, label + ": ", size=11, bold=True, color=NAVY)
    add_run(p, value, size=11, color=LIGHT_BLUE)


# ----- PAGE 2 -----
doc.add_page_break()
banner()
d_title("Professional Profile")
d_para("[ Write a short introductory paragraph here: your profile, years of "
       "experience, key strengths, and your career objective. ]")
d_title("Areas of Expertise")
for _ in range(5):
    d_bullet("[ Area of expertise / core skill ]")
d_title("Key Projects")
d_entry("[ Project name ]", "[ Company - City ]", "[ Dates ]")
d_bullet("[ Project description / result ]")
d_bullet("[ Project description / result ]")
d_entry("[ Project name ]", "[ Company - City ]", "[ Dates ]")
d_bullet("[ Project description / result ]")
d_bullet("[ Project description / result ]")

# ----- PAGE 3 -----
doc.add_page_break()
banner()
d_title("Certifications")
for _ in range(3):
    d_bullet("[ Certification - Issuing body - Year ]")
d_title("Additional Training")
d_entry("[ Course / Workshop ]", "[ Institution - City ]", "[ Year ]")
d_bullet("[ Content / skills acquired ]")
d_entry("[ Course / Workshop ]", "[ Institution - City ]", "[ Year ]")
d_bullet("[ Content / skills acquired ]")
d_title("Key Achievements")
for _ in range(3):
    d_bullet("[ Achievement or measurable result ]")
d_title("References")
d_label("[ Reference name ]", "[ Title, Company ]")
d_label("Contact", "[ Phone · email ]")
d_para(" ")
d_label("[ Reference name ]", "[ Title, Company ]")
d_label("Contact", "[ Phone · email ]")

doc.save(OUT)
print("Saved", OUT, "(3-page blank English template)")
